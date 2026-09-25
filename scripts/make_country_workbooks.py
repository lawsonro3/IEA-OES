#!/usr/bin/env python3
"""Make the yearly data-collection workbooks for IEA-OES country representatives.

One short workbook per country, pre-filled from the lists: the country's test sites, and
the projects its chapters have named since --since. Reps update those rows and add new
ones. Each data sheet opens with one example row, marked as an example. Also writes a
blank template. The report year is in the file name only:
"IEA OES Annual Test Site and Project Reporting - <Country> - <Year>.xlsx".

    python3 scripts/make_country_workbooks.py --year 2026                 # blank + every country
    python3 scripts/make_country_workbooks.py --year 2026 --country Portugal

Run extract.py first (it writes output/project_years.csv, used to pick recent projects).
"""

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent.parent
OUT = HERE / "output" / "country_workbooks"
ROWS = 100      # input rows prepared on each sheet
HEAD = 3        # header row; data starts on the next row

FONT = "Arial"
F_INPUT = PatternFill("solid", fgColor="FFF6CC")     # yellow: fill in
F_EX = PatternFill("solid", fgColor="EEF1F4")        # grey: the example row
F_HEAD = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="C9D3DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LISTS = {
    "site_type": ["open water", "laboratory", "shore-based", "mixed"],
    "site_status": ["operational", "under development", "closed"],
    "technology": ["wave", "tidal stream", "tidal range", "river current", "ocean current", "OTEC",
                   "salinity gradient", "other"],
    "project_status": ["planned", "under construction", "operating", "removed or cancelled"],
}
SITE_TYPE = {"open-water": "open water"}

# (header, width, dropdown, note shown when hovering over the heading)
SITE_COLS = [
    ("Site name", 38, None, None),
    ("Type", 14, "site_type", "open water: sea, estuary or river. laboratory: tank, basin, flume or test rig. "
                              "shore-based: plant on a breakwater or coast. mixed: several of these."),
    ("Location", 30, None, "Town, region or water body."),
    ("Latitude", 11, None, "Optional. Decimal degrees, e.g. 58.96."),
    ("Longitude", 11, None, "Optional. Decimal degrees, e.g. -3.30 (west is negative)."),
    ("Status", 18, "site_status", "At 31 December."),
    ("Website", 32, None, None),
    ("Notes", 40, None, None),
]
PROJECT_COLS = [
    ("Project name", 38, None, None),
    ("Developer", 26, None, None),
    ("Technology", 15, "technology", None),
    ("Location or test site", 32, None, "The test site's name if it is at one; otherwise town, region or water body."),
    ("Latitude", 11, None, "Optional. Decimal degrees."),
    ("Longitude", 11, None, "Optional. Decimal degrees (west is negative)."),
    ("Status", 20, "project_status", "At 31 December."),
    ("Capacity (kW)", 13, None, "Installed capacity if operating; planned capacity otherwise. 1 MW = 1000 kW."),
    ("Website", 32, None, None),
    ("Notes", 40, None, None),
]

EXAMPLE_NOTE = "EXAMPLE - delete this row"
EXAMPLE_SITE = ["European Marine Energy Centre (EMEC)", "open water", "Orkney, Scotland", 58.96, -3.30,
                "operational", "https://www.emec.org.uk/", EXAMPLE_NOTE]
EXAMPLE_PROJECT = ["Shetland Tidal Array", "Nova Innovation", "tidal stream", "Bluemull Sound, Shetland",
                   60.70, -0.98, "operating", 600, "https://www.novainnovation.com/", EXAMPLE_NOTE]


def split(s, sep=" | "):
    return [x.strip() for x in (s or "").split(sep) if x.strip()]


def font(**kw):
    return Font(name=FONT, **kw)


def instructions(wb, country, year, since):
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80
    rows = [
        ("Country", country or ""),
        ("Your name", ""),
        ("Email", ""),
        (None, None),
        ("What to do", ""),
        ("Test sites", "Fill out information for the test sites you would like included in the IEA OES "
                       "project map for this reporting year. Laboratories and tanks count."),
        ("Projects", "Fill out information for the projects you would like included in the IEA OES "
                     "project map for this reporting year, including planned projects."),
        ("Tips", "Fill in the yellow cells. Hover over a column heading for help."),
    ]
    for i, (a, b) in enumerate(rows, 1):
        if a is None:
            continue
        ca = ws.cell(row=i, column=1, value=a)
        ca.font = font(bold=True, size=11 if a == "What to do" else 10,
                       color="1F4E78" if a == "What to do" else "000000")
        cb = ws.cell(row=i, column=2, value=b or None)
        cb.font = font(size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if a in ("Country", "Your name", "Email"):
            cb.fill = F_INPUT
            cb.border = BORDER


def data_sheet(wb, name, cols, rows, refs, title, example):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = font(bold=True, size=13, color="1F4E78")
    for j, (head, width, lst, note) in enumerate(cols, 1):
        c = ws.cell(row=HEAD, column=j, value=head)
        c.font = font(bold=True, color="FFFFFF", size=10)
        c.fill = F_HEAD
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if note:
            c.comment = Comment(note, "IEA-OES")
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[HEAD].height = 30
    rows = [example] + rows
    for i in range(ROWS):
        vals = rows[i] if i < len(rows) else None
        for j in range(1, len(cols) + 1):
            v = vals[j - 1] if vals else None
            c = ws.cell(row=HEAD + 1 + i, column=j, value=v if v != "" else None)
            c.border = BORDER
            if i == 0:      # the example row
                c.font = font(size=10, italic=True, color="5B6B78", bold=j == len(cols))
                c.fill = F_EX
            else:
                c.font = font(size=10)
                c.fill = F_INPUT
    end = HEAD + ROWS
    for j, (head, width, lst, note) in enumerate(cols, 1):
        if lst:
            dv = DataValidation(type="list", formula1=refs[lst], allow_blank=True,
                                showErrorMessage=True, errorStyle="warning",
                                error="Please pick from the list, or explain in Notes.")
            dv.add(f"{get_column_letter(j)}{HEAD + 1}:{get_column_letter(j)}{end}")
            ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(row=HEAD + 1, column=2)


def lists_sheet(wb):
    ws = wb.create_sheet("Lists")
    refs = {}
    for j, (key, values) in enumerate(LISTS.items(), 1):
        ws.cell(row=1, column=j, value=key).font = font(bold=True, size=10)
        for i, v in enumerate(values, 2):
            ws.cell(row=i, column=j, value=v).font = font(size=10)
        col = get_column_letter(j)
        refs[key] = f"Lists!${col}$2:${col}${len(values) + 1}"
    ws.sheet_state = "hidden"
    return refs


def load(country, since):
    sites = list(csv.DictReader(open(HERE / "gazetteers" / "test_sites.csv", encoding="utf-8")))
    projects = list(csv.DictReader(open(HERE / "gazetteers" / "projects.csv", encoding="utf-8")))
    years = defaultdict(set)
    py = HERE / "output" / "project_years.csv"
    if py.exists():
        for r in csv.DictReader(open(py, encoding="utf-8")):
            if any(c != "(outside country chapters)" for c in split(r["reporting_countries"], "; ")):
                years[r["project_id"]].add(int(r["report_year"]))
    for p in projects:
        if p["report_years_basis"].startswith("Pilot"):
            years[p["project_id"]].update(int(y) for y in split(p["report_years"], "; "))
    site_name = {s["site_id"]: s["name"] for s in sites}

    site_rows = sorted(
        [[s["name"], SITE_TYPE.get(s["site_type"], s["site_type"]), s["location"],
          float(s["lat"]) if s["lat"] else "", float(s["lon"]) if s["lon"] else "", "",
          (split(s["primre_website"]) or [""])[0], ""]
         for s in sites if s["country"] == country], key=lambda r: r[0])

    project_rows = []
    for p in projects:
        chapters = split(p["reporting_country"]) or [p["country"]]
        ys = years.get(p["project_id"])
        if country not in chapters or p["in_scope"] != "yes" or not ys or max(ys) < since:
            continue
        loc = site_name.get(p["test_site_id"]) or p["location"].strip()
        if len(loc) > 60 or loc[:1].islower():      # a sentence from a report, not a place
            loc = ""
        tech = p["technology"] if p["technology"] in LISTS["technology"] else \
            {"tidal": "tidal stream"}.get(p["technology"], "other")
        project_rows.append([p["name"], p["developer"], tech, loc,
                             float(p["lat"]) if p["lat"] else "", float(p["lon"]) if p["lon"] else "",
                             "", "", (split(p["primre_website"]) or [""])[0], ""])
    project_rows.sort(key=lambda r: r[0].lower())
    return site_rows, project_rows


def build(country, year, since, path):
    wb = Workbook()
    instructions(wb, country, year, since)
    refs = lists_sheet(wb)
    site_rows, project_rows = load(country, since) if country else ([], [])
    who = f": {country}" if country else ""
    data_sheet(wb, "Test sites", SITE_COLS, site_rows, refs, f"Test sites{who}", EXAMPLE_SITE)
    data_sheet(wb, "Projects", PROJECT_COLS, project_rows, refs, f"Projects{who}", EXAMPLE_PROJECT)
    wb.move_sheet("Lists", offset=10)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(site_rows), len(project_rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--year", type=int, required=True, help="the report year the return covers")
    ap.add_argument("--country", help="one country (default: every country with report chapters)")
    ap.add_argument("--since", type=int, default=2020, help="pre-fill projects named in this year or later")
    args = ap.parse_args()

    for old in OUT.glob("*.xlsx"):
        old.unlink()
    blank = OUT / "TEMPLATE - IEA OES Annual Test Site and Project Reporting - Country Name - Year.xlsx"
    build(None, args.year, args.since, blank)
    print(f"Wrote {blank.relative_to(HERE)}")
    if args.country:
        countries = [args.country]
    else:
        projects = csv.DictReader(open(HERE / "gazetteers" / "projects.csv", encoding="utf-8"))
        countries = sorted({c for p in projects for c in split(p["reporting_country"])})
    for c in countries:
        path = OUT / f"IEA OES Annual Test Site and Project Reporting - {c} - {args.year}.xlsx"
        n_s, n_p = build(c, args.year, args.since, path)
        print(f"Wrote {path.relative_to(HERE)}: {n_s} test sites, {n_p} projects pre-filled")


if __name__ == "__main__":
    main()
