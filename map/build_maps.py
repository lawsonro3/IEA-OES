#!/usr/bin/env python3
"""Build the two map pages from the lists and the extraction workbook.

  output/test_sites_map.html  OES Test Sites Atlas: every test site the reports name,
                              by site type, with the projects linked to each site
  output/projects_map.html    Ocean Energy Projects Atlas: one marker per project that
                              has a stated capacity and is named in a country chapter

Each page is self-contained (data, Leaflet stylesheet and Natural Earth outlines are
embedded); only the Leaflet and topojson-client scripts load from a CDN.

Run extract.py and scripts/add_map_fields.py first, then:
    python3 map/build_maps.py
"""

import csv
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

MAP = Path(__file__).resolve().parent
HERE = MAP.parent
OUT = HERE / "output"
WORKBOOK = OUT / "oes_test_sites_and_projects.xlsx"
PILOT = HERE.parent / "Extraction Pilot" / "tables"
PRIMRE = HERE.parent / "PRIMRE databases" / "primre-projects.csv"
OUTSIDE = "(outside country chapters)"
YEARS = list(range(2002, 2026))

TECH = {"wave": "wave", "tidal stream": "tidal", "tidal range": "tidal", "tidal": "tidal",
        "river current": "river current", "ocean current": "ocean current", "OTEC": "OTEC",
        "salinity gradient": "salinity gradient"}
PREC = {"test site": "test site", "PRIMRE": "primre", "place": "place", "region": "region",
        "country": "country"}


def read_csv(path, enc="utf-8"):
    with open(path, newline="", encoding=enc) as f:
        return list(csv.DictReader(f))


def read_sheet(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    header = next(rows)
    return [dict(zip(header, r)) for r in rows]


def split(s, sep=" | "):
    return [x.strip() for x in (s or "").split(sep) if x.strip()]


def clip(text, n=420):
    text = re.sub(r"\s+", " ", (text or "")).strip()
    return text if len(text) <= n else text[:n - 3].rsplit(" ", 1)[0] + "…"


def year_ranges(years):
    """[2016, 2017, 2018, 2020] -> '2016–18, 2020'."""
    ys, out, i = sorted(set(years)), [], 0
    while i < len(ys):
        j = i
        while j + 1 < len(ys) and ys[j + 1] == ys[j] + 1:
            j += 1
        out.append(str(ys[i]) if i == j else f"{ys[i]}–{str(ys[j])[2:]}")
        i = j + 1
    return ", ".join(out)


def build_pages(template, data, placeholder):
    html = (template
            .replace("/*__LEAFLET_CSS__*/", (MAP / "leaflet.css").read_text(encoding="utf-8"))
            .replace(placeholder, json.dumps(data, ensure_ascii=False, separators=(",", ":")))
            .replace("__COUNTRIES_TOPO__", (MAP / "countries-10m.json").read_text(encoding="utf-8")))
    return html


def main():
    wb = load_workbook(WORKBOOK, read_only=True)
    sites = read_csv(HERE / "gazetteers" / "test_sites.csv")
    projects = read_csv(HERE / "gazetteers" / "projects.csv")
    pilot_obs = {o["obs_id"]: o for o in read_csv(PILOT / "observations.csv")}
    primre = {r["Name"]: r for r in read_csv(PRIMRE, "utf-8-sig")}
    site_by_id = {s["site_id"]: s for s in sites}

    # ---- Project years, country chapters only: extraction mentions plus the Pilot's years.
    chap_n = defaultdict(lambda: defaultdict(int))
    chap_ex = {}
    for m in read_sheet(wb, "Project mentions"):
        if not m["project_id"] or m["reporting_country"] == OUTSIDE:
            continue
        pid, y = m["project_id"], int(m["report_year"])
        chap_n[pid][y] += 1
        best = chap_ex.get((pid, y))
        if best is None or (m["resolution"] == "unique term" and best[2] != "unique term"):
            chap_ex[(pid, y)] = (clip(m["context"]), f"{m['reporting_country']} chapter · line {m['line']}",
                                 m["resolution"])
    pilot_ex = {}
    for p in projects:
        for oid in split(p["obs_ids"], "; "):
            o = pilot_obs.get(oid)
            if not o:
                continue
            y = int(o["report_year"])
            pilot_ex.setdefault((p["project_id"], y), (
                clip(o["evidence"]), f"{o['reporting_country']} chapter · p. {o['printed_page']} (Pilot)"))
            chap_n[p["project_id"]][y] = max(chap_n[p["project_id"]][y], 1)

    # ---- Test sites
    mentions = defaultdict(lambda: defaultdict(int))
    site_ex = {}
    for m in read_sheet(wb, "Test site mentions"):
        sid, y = m["site_id"], int(m["report_year"])
        mentions[sid][y] += 1
        site_ex.setdefault((sid, y), [clip(m["context"]), m["reporting_country"], m["line"]])
    at_site = defaultdict(list)
    for p in projects:
        if p["test_site_id"] and chap_n.get(p["project_id"]):
            ys = sorted(chap_n[p["project_id"]])
            at_site[p["test_site_id"]].append((ys[0], p["name"], year_ranges(ys)))
    site_rows = []
    for s in sites:
        ys = sorted(mentions[s["site_id"]])
        site_rows.append({
            "id": s["site_id"], "name": s["name"], "country": s["country"], "location": s["location"],
            "lat": float(s["lat"]), "lon": float(s["lon"]),
            "approx": s["coord_note"].startswith("approximate") or not s["coord_note"].startswith("stated"),
            "type": s["site_type"],
            "years": {str(y): mentions[s["site_id"]][y] for y in ys},
            "ex": {str(y): site_ex[(s["site_id"], y)] for y in ys},
            "projects": [[n, r] for _, n, r in sorted(at_site[s["site_id"]])],
        })
    sites_data = {"sites": site_rows, "years": YEARS}

    # ---- Projects: ocean energy, a stated capacity, named in a country chapter.
    rows, no_capacity, no_location = [], 0, 0
    for p in projects:
        pid = p["project_id"]
        if p["in_scope"] != "yes" or not chap_n.get(pid):
            continue
        if not p["capacity_kw"]:
            no_capacity += 1
            continue
        lat = float(p["lat"]) if p["lat"] else None
        lon = float(p["lon"]) if p["lon"] else None
        prec = PREC.get(p["location_basis"], "country")
        if lat is None and not p["country"]:
            no_location += 1
            continue
        site = site_by_id.get(p["test_site_id"])
        precs = [primre[n] for n in split(p["primre_name"]) if n in primre]
        if prec == "test site" and site:
            where = site["name"]
        elif prec == "primre" and precs:
            where = precs[0]["Waterbody"] or p["country"]
        elif prec in ("place", "region"):
            where = clip(p["location"], 70) or p["country"]
        else:
            where = p["country"]
        ys = sorted(chap_n[pid])
        ex = {}
        for y in ys:
            e = pilot_ex.get((pid, y)) or (chap_ex[(pid, y)][:2] if (pid, y) in chap_ex else None)
            if e:
                ex[str(y)] = list(e)
        pilot_years = [int(y) for y in split(p["report_years"], "; ")] \
            if p["report_years_basis"].startswith("Pilot") else []
        status, status_year = p["latest_status"], str(max(pilot_years)) if pilot_years else ""
        if not status and precs:
            status, status_year = precs[0]["Project Status"], "PRIMRE"
        rows.append({
            "id": pid, "name": p["name"], "kw": float(p["capacity_kw"]),
            "capBasis": "the report" if p["capacity_basis"].startswith("Pilot") else "PRIMRE",
            "tech": TECH.get(p["technology"], "hybrid / multiple"), "techRaw": p["technology"],
            "status": status, "statusYear": status_year,
            "country": p["country"], "lat": lat, "lon": lon, "prec": prec, "where": where,
            "site": site["name"] if site else "", "dev": p["developer"],
            "years": {str(y): chap_n[pid][y] for y in ys}, "ex": ex,
            "primre": [[n, primre[n]["Website"]] for n in split(p["primre_name"]) if n in primre],
        })
    projects_data = {"projects": rows, "years": YEARS,
                     "excluded": {"no_capacity": no_capacity, "no_location": no_location}}

    OUT.mkdir(exist_ok=True)
    (OUT / "test_sites_map.html").write_text(build_pages(
        (MAP / "test_sites_template.html").read_text(encoding="utf-8"), sites_data, "__SITE_DATA__"),
        encoding="utf-8")
    shared_css = (MAP / "test_sites_template.html").read_text(encoding="utf-8") \
        .split("<style>")[2].split("</style>")[0]
    (OUT / "projects_map.html").write_text(build_pages(
        (MAP / "projects_template.html").read_text(encoding="utf-8").replace("/*__SHARED_CSS__*/", shared_css),
        projects_data, "__PROJECT_DATA__"), encoding="utf-8")

    named = sum(1 for s in site_rows if s["years"])
    print(f"Test sites: {named} of {len(site_rows)} named in the reports; "
          f"{sum(len(s['projects']) for s in site_rows)} project links")
    by = defaultdict(int)
    for r in rows:
        by[r["prec"]] += 1
    print(f"Projects: {len(rows)} mapped ({', '.join(f'{v} {k}' for k, v in sorted(by.items()))}); "
          f"not shown: {no_capacity} without capacity, {no_location} without location")
    for f in ("test_sites_map.html", "projects_map.html"):
        print(f"Wrote output/{f} ({(OUT / f).stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
